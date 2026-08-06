#!/usr/bin/env python3
"""
CoLab Space Point — Premium Corporate Digital Agency Proposal
Dark Blue & White Corporate Theme | Microsoft Word (.docx)
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor, Twips, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import copy

# ── Brand Colors ──────────────────────────────────────────────
NAVY = RGBColor(0x0A, 0x25, 0x40)       # Deep navy
NAVY_MID = RGBColor(0x1B, 0x3A, 0x5F)   # Mid navy
ACCENT = RGBColor(0x1E, 0x5F, 0xA8)     # Corporate blue
LIGHT_BLUE = RGBColor(0xE8, 0xF0, 0xF8) # Soft blue bg
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x2E)
GRAY = RGBColor(0x5A, 0x5A, 0x6E)
LIGHT_GRAY = RGBColor(0xF5, 0xF7, 0xFA)
GOLD = RGBColor(0xC9, 0xA2, 0x27)       # Subtle gold accent
BORDER_BLUE = "1B3A5F"
HEADER_FILL = "0A2540"
ACCENT_FILL = "1E5FA8"
ROW_ALT = "E8F0F8"
ROW_WHITE = "FFFFFF"
CALLOUT_FILL = "F0F5FA"
LIGHT_FILL = "F5F7FA"

OUTPUT = "/workspace/proposals/CoLab_Space_Point_Digital_Agency_Proposal.docx"


# ── Helpers ───────────────────────────────────────────────────

def set_run(run, *, size=11, bold=False, color=DARK, font="Calibri", italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = font
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:eastAsia"), font)
    rFonts.set(qn("w:cs"), font)


def add_para(doc, text="", *, size=11, bold=False, color=DARK, align="left",
             space_before=0, space_after=6, font="Calibri", italic=False):
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    if text:
        run = p.add_run(text)
        set_run(run, size=size, bold=bold, color=color, font=font, italic=italic)
    return p


def add_rich_para(doc, parts, *, align="left", space_before=0, space_after=6):
    """parts = list of (text, kwargs)"""
    p = doc.add_paragraph()
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    }[align]
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.15
    for text, kwargs in parts:
        run = p.add_run(text)
        set_run(run, **kwargs)
    return p


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_borders(cell, color=BORDER_BLUE, sz="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)


def clear_cell_margins(cell, top=40, bottom=40, left=80, right=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(mar)


def cell_text(cell, text, *, size=10, bold=False, color=DARK, align="left", font="Calibri"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, font=font)


def add_page_break(doc):
    doc.add_page_break()


def add_horizontal_line(doc, color=HEADER_FILL):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="18" w:space="1" w:color="{color}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
    return p


def add_thin_line(doc, color="1E5FA8"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="{color}"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
    return p


def section_header(doc, title, subtitle=None):
    add_para(doc, title, size=26, bold=True, color=NAVY, font="Georgia",
             space_before=6, space_after=4)
    add_horizontal_line(doc)
    if subtitle:
        add_para(doc, subtitle, size=11, color=GRAY, italic=True,
                 space_before=0, space_after=14)


def subsection(doc, title):
    add_para(doc, title, size=14, bold=True, color=NAVY_MID, font="Calibri",
             space_before=14, space_after=6)


def body(doc, text, *, space_after=8):
    add_para(doc, text, size=11, color=DARK, align="justify", space_after=space_after)


def bullet(doc, text, *, indent=0.25):
    p = add_para(doc, f"●  {text}", size=11, color=DARK, space_before=1, space_after=3)
    p.paragraph_format.left_indent = Inches(indent)
    return p


def check_item(doc, text):
    p = add_para(doc, f"✓  {text}", size=10.5, color=DARK, space_before=1, space_after=2)
    p.paragraph_format.left_indent = Inches(0.2)
    return p


def callout_box(doc, title, text, fill=CALLOUT_FILL):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_borders(cell, color="1E5FA8", sz="12")
    clear_cell_margins(cell, top=80, bottom=80, left=120, right=120)
    cell.text = ""
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(title)
    set_run(r1, size=11, bold=True, color=NAVY)
    p1.paragraph_format.space_after = Pt(4)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(text)
    set_run(r2, size=10.5, color=DARK)
    p2.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def icon_row(doc, items):
    """items = list of (icon_label, title, description)"""
    cols = len(items)
    table = doc.add_table(rows=1, cols=cols)
    table.autofit = True
    for i, (icon, title, desc) in enumerate(items):
        cell = table.cell(0, i)
        shade_cell(cell, LIGHT_FILL)
        set_cell_borders(cell, color="D0D8E0", sz="4")
        clear_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        cell.text = ""
        p0 = cell.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(icon)
        set_run(r0, size=16, bold=True, color=ACCENT)
        p0.paragraph_format.space_after = Pt(4)
        p1 = cell.add_paragraph()
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(title)
        set_run(r1, size=10, bold=True, color=NAVY)
        p1.paragraph_format.space_after = Pt(3)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(desc)
        set_run(r2, size=8.5, color=GRAY)
        p2.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def styled_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.autofit = True
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, HEADER_FILL)
        set_cell_borders(cell, color=HEADER_FILL, sz="4")
        clear_cell_margins(cell)
        cell_text(cell, h, size=10, bold=True, color=WHITE, align="center")

    for r_idx, row in enumerate(rows):
        fill = ROW_ALT if r_idx % 2 == 0 else ROW_WHITE
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            shade_cell(cell, fill)
            set_cell_borders(cell, color="C8D0D8", sz="4")
            clear_cell_margins(cell)
            align = "center" if c_idx > 0 else "left"
            bold = c_idx == 0
            cell_text(cell, val, size=9.5, bold=bold, color=DARK, align=align)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    return table


def pricing_header_card(doc, name, price, suitable, highlight=False):
    fill = HEADER_FILL if highlight else ACCENT_FILL
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_borders(cell, color=fill, sz="4")
    clear_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    cell.text = ""
    p1 = cell.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run(name.upper())
    set_run(r1, size=14, bold=True, color=WHITE, font="Georgia")
    p1.paragraph_format.space_after = Pt(6)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(price)
    set_run(r2, size=22, bold=True, color=WHITE, font="Georgia")
    p2.paragraph_format.space_after = Pt(4)
    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(suitable)
    set_run(r3, size=10, color=RGBColor(0xD0, 0xE0, 0xF0), italic=True)
    p3.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)


def three_col_packages(doc, packages):
    """packages = list of 3 dicts with name, price, suitable, features"""
    table = doc.add_table(rows=2, cols=3)
    table.autofit = True

    for i, pkg in enumerate(packages):
        # Header
        cell = table.rows[0].cells[i]
        fill = HEADER_FILL if i == 1 else ACCENT_FILL
        shade_cell(cell, fill)
        set_cell_borders(cell, color=fill, sz="4")
        clear_cell_margins(cell, top=80, bottom=80, left=60, right=60)
        cell.text = ""
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(pkg["name"].upper())
        set_run(r1, size=11, bold=True, color=WHITE, font="Georgia")
        p1.paragraph_format.space_after = Pt(4)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(pkg["price"])
        set_run(r2, size=16, bold=True, color=WHITE)
        p2.paragraph_format.space_after = Pt(3)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(pkg["suitable"])
        set_run(r3, size=8, color=RGBColor(0xC8, 0xD8, 0xE8), italic=True)

        # Features
        cell2 = table.rows[1].cells[i]
        shade_cell(cell2, LIGHT_FILL if i != 1 else ROW_ALT)
        set_cell_borders(cell2, color="C8D0D8", sz="4")
        clear_cell_margins(cell2, top=60, bottom=60, left=60, right=60)
        cell2.text = ""
        first = True
        for feat in pkg["features"]:
            if first:
                p = cell2.paragraphs[0]
                first = False
            else:
                p = cell2.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(f"✓  {feat}")
            set_run(r, size=8.5, color=DARK)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def set_narrow_margins(section):
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)


def add_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CoLab Space Point  |  Digital Agency Proposal  |  Confidential  |  www.colabpoint.com")
    set_run(run, size=8, color=GRAY)


def add_header(doc):
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("COLAB SPACE POINT")
    set_run(run, size=8, bold=True, color=NAVY)


# ── Document Sections ─────────────────────────────────────────

def build_cover(doc):
    # Top accent bar via table
    bar = doc.add_table(rows=1, cols=1)
    cell = bar.cell(0, 0)
    shade_cell(cell, HEADER_FILL)
    set_cell_borders(cell, color=HEADER_FILL, sz="2")
    clear_cell_margins(cell, top=50, bottom=50, left=40, right=40)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PREMIUM DIGITAL AGENCY PROPOSAL")
    set_run(r, size=9, bold=True, color=WHITE)
    p.paragraph_format.space_after = Pt(2)

    add_para(doc, "", space_after=28)
    add_para(doc, "COLAB SPACE POINT", size=34, bold=True, color=NAVY,
             align="center", font="Georgia", space_after=6)
    add_para(doc, "Digital Agency Proposal", size=18, color=ACCENT,
             align="center", font="Georgia", space_after=10)

    add_thin_line(doc)

    add_para(doc, "Website Designing & Development  ·  Digital Marketing  ·  Social Media Management",
             size=10.5, color=GRAY, align="center", space_before=6, space_after=16)

    # Services highlight boxes
    svc_table = doc.add_table(rows=1, cols=3)
    services_cover = [
        ("◆", "WEBSITE DESIGN\n& DEVELOPMENT"),
        ("◆", "DIGITAL\nMARKETING"),
        ("◆", "SOCIAL MEDIA\nMANAGEMENT"),
    ]
    for i, (icon, label) in enumerate(services_cover):
        cell = svc_table.cell(0, i)
        shade_cell(cell, LIGHT_FILL)
        set_cell_borders(cell, color="1B3A5F", sz="8")
        clear_cell_margins(cell, top=80, bottom=80, left=60, right=60)
        cell.text = ""
        p1 = cell.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(icon)
        set_run(r1, size=14, color=ACCENT)
        p1.paragraph_format.space_after = Pt(4)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(label)
        set_run(r2, size=9, bold=True, color=NAVY)

    add_para(doc, "", space_after=36)

    # Bottom info
    add_thin_line(doc)
    add_para(doc, "Prepared by CoLab Space Point", size=11, bold=True, color=NAVY,
             align="center", space_before=8, space_after=3)
    add_para(doc, "2nd Floor, Anwar Center, Madina Road Near Gymkhana, Gujrat, Pakistan",
             size=10, color=GRAY, align="center", space_after=2)
    add_para(doc, "+92 349 7684322  ·  +92 332 4384322  ·  www.colabpoint.com",
             size=10, color=GRAY, align="center", space_after=2)
    add_para(doc, "colabpoint@gmail.com  ·  hello@colabpoint.com",
             size=10, color=GRAY, align="center", space_after=8)

    # Bottom navy bar
    bar2 = doc.add_table(rows=1, cols=1)
    cell = bar2.cell(0, 0)
    shade_cell(cell, HEADER_FILL)
    set_cell_borders(cell, color=HEADER_FILL, sz="2")
    clear_cell_margins(cell, top=30, bottom=30, left=40, right=40)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CONFIDENTIAL — FOR CLIENT REVIEW ONLY")
    set_run(r, size=8, bold=True, color=WHITE)

    add_page_break(doc)


def build_toc(doc):
    section_header(doc, "Table of Contents",
                   "A structured overview of this proposal")

    toc_items = [
        ("00", "Executive Summary", "Opportunity overview and proposal highlights"),
        ("01", "Company Profile", "Who we are and what we deliver"),
        ("02", "Why Choose CoLab Space Point", "The advantages of partnering with us"),
        ("03", "Our Services & Investment Overview", "Portfolio summary with package pricing"),
        ("04", "Website Designing & Development", "Packages, features & comparison"),
        ("05", "Social Media Management", "Monthly engagement packages"),
        ("06", "Digital Marketing", "Performance-driven growth packages"),
        ("07", "Add-On Services", "Optional premium enhancements"),
        ("08", "Our Process", "From discovery to long-term support"),
        ("09", "Why Clients Trust Us", "Partnership philosophy & commitment"),
        ("10", "Frequently Asked Questions", "Answers to common client queries"),
        ("11", "Contact & Next Steps", "How to engage with our team"),
    ]

    table = doc.add_table(rows=len(toc_items), cols=3)
    for i, (num, title, desc) in enumerate(toc_items):
        c0 = table.rows[i].cells[0]
        c1 = table.rows[i].cells[1]
        c2 = table.rows[i].cells[2]
        fill = ROW_ALT if i % 2 == 0 else ROW_WHITE
        for c in (c0, c1, c2):
            shade_cell(c, fill)
            set_cell_borders(c, color="E0E4E8", sz="2")
            clear_cell_margins(c, top=50, bottom=50, left=60, right=40)
        cell_text(c0, num, size=14, bold=True, color=ACCENT, align="center")
        cell_text(c1, title, size=11, bold=True, color=NAVY)
        cell_text(c2, desc, size=9, color=GRAY)
        c0.width = Inches(0.7)
        c1.width = Inches(3.5)
        c2.width = Inches(2.8)

    add_page_break(doc)


def build_company_profile(doc):
    section_header(doc, "01  |  Company Profile",
                   "Building digital excellence from the heart of Gujrat")

    body(doc,
         "CoLab Space Point is a premier digital solutions partner based in Gujrat, Pakistan, "
         "operating under the trusted CoLab Point ecosystem established in 2021. We specialize "
         "in delivering end-to-end digital business solutions that empower startups, growing "
         "enterprises, and established brands to compete and thrive in the modern digital economy. "
         "Rooted in collaboration, innovation, and measurable results, we combine strategic consulting, "
         "creative design excellence, and performance-driven marketing to accelerate growth.")

    callout_box(doc, "OUR MISSION",
                "To provide complete digital business solutions under one roof — enabling entrepreneurs, "
                "startups, and enterprises to build strong digital foundations, expand market reach, "
                "and achieve sustainable growth through modern technology and strategic expertise.")

    subsection(doc, "Complete Digital Business Solutions")

    services = [
        ("Website Designing & Development", "Conversion-focused websites for performance & scale"),
        ("WordPress Development", "CMS-powered sites with flexible content management"),
        ("E-Commerce Solutions", "Online stores with payment gateways & inventory tools"),
        ("Branding & Identity", "Brand systems that communicate authority and trust"),
        ("Social Media Management", "Consistent presence across major social platforms"),
        ("Digital Marketing", "Campaigns engineered for leads, sales, and ROI"),
        ("Search Engine Optimization (SEO)", "Organic visibility strategies that compound over time"),
        ("Google Ads Management", "High-intent paid search with disciplined budgets"),
        ("Meta Ads Management", "Facebook & Instagram ads for awareness & conversion"),
        ("Business Growth Solutions", "Holistic strategies aligned to commercial objectives"),
    ]

    table = doc.add_table(rows=len(services), cols=2)
    for i, (svc, desc) in enumerate(services):
        c0, c1 = table.rows[i].cells
        fill = ROW_ALT if i % 2 == 0 else ROW_WHITE
        for c in (c0, c1):
            shade_cell(c, fill)
            set_cell_borders(c, color="D0D8E0", sz="2")
            clear_cell_margins(c, top=24, bottom=24, left=60, right=40)
        cell_text(c0, f"▸  {svc}", size=9.5, bold=True, color=NAVY)
        cell_text(c1, desc, size=9, color=DARK)
        c0.width = Inches(2.8)
        c1.width = Inches(4.2)

    add_para(doc, "", space_after=4)
    body(doc,
         "Operating from Anwar Center, Madina Road, Gujrat, CoLab Point has built a reputation as a hub "
         "for freelancers, startups, and digital professionals. CoLab Space Point extends this legacy into "
         "full-service digital agency capabilities — combining local market understanding with "
         "international-standard delivery quality.")

    add_page_break(doc)


def build_why_choose(doc):
    section_header(doc, "02  |  Why Choose CoLab Space Point",
                   "Eight pillars that define our partnership standard")

    body(doc,
         "Selecting a digital partner is a strategic decision. CoLab Space Point is structured to "
         "deliver not merely services, but a lasting competitive advantage for your business. "
         "The following pillars define how we work and why clients choose to grow with us.")

    pillars = [
        ("01", "Experienced Team",
         "Proven expertise across web development, design, SEO, paid media, and brand strategy — "
         "professionals who understand both technology and commercial outcomes."),
        ("02", "Professional Support",
         "Dedicated support from onboarding through launch and beyond. We respond promptly, "
         "communicate proactively, and resolve issues before they become obstacles."),
        ("03", "Business-Focused Solutions",
         "Every recommendation is anchored in your business objectives — solutions designed to "
         "generate inquiries, sales, and brand equity, not generic templates."),
        ("04", "Modern Technologies",
         "WordPress, WooCommerce, analytics, and advertising platforms — secure, scalable, "
         "and future-ready digital assets built on best-practice frameworks."),
        ("05", "Creative Design",
         "Visual excellence balanced with usability — digital experiences that reflect your "
         "brand's premium positioning."),
        ("06", "Performance Marketing",
         "Campaigns engineered around measurable KPIs — cost per lead, ROAS, engagement, and "
         "conversion quality — optimized so your investment works harder."),
        ("07", "Transparent Communication",
         "Clear scopes, honest timelines, structured reporting, and open dialogue. You always "
         "know what is delivered and why it matters."),
        ("08", "Long-Term Partnership",
         "Success measured by relationship longevity — ongoing optimization, support, and "
         "strategic evolution, not one-off transactions."),
    ]

    # 2-column compact pillar grid (4 rows)
    grid = doc.add_table(rows=4, cols=2)
    for i, (num, title, desc) in enumerate(pillars):
        cell = grid.rows[i // 2].cells[i % 2]
        shade_cell(cell, LIGHT_FILL)
        set_cell_borders(cell, color="1B3A5F", sz="6")
        clear_cell_margins(cell, top=40, bottom=40, left=60, right=50)
        cell.text = ""
        p1 = cell.paragraphs[0]
        r1 = p1.add_run(f"{num}  {title}")
        set_run(r1, size=10, bold=True, color=NAVY)
        p1.paragraph_format.space_after = Pt(3)
        p2 = cell.add_paragraph()
        r2 = p2.add_run(desc)
        set_run(r2, size=8.5, color=DARK)
        p2.paragraph_format.space_after = Pt(1)

    add_page_break(doc)


def build_services_overview(doc):
    section_header(doc, "03  |  Our Services Overview",
                   "Integrated capabilities for complete digital growth")

    body(doc,
         "CoLab Space Point operates as a full-service digital agency. Whether you require a "
         "high-performance website, a structured social media presence, or a results-oriented "
         "advertising program, our teams deliver with consistency, accountability, and strategic depth.")

    icon_row(doc, [
        ("◆ WEB", "Website Solutions", "Design, WordPress & E-Commerce"),
        ("◆ SOCIAL", "Social Media", "Content, Community & Growth"),
        ("◆ ADS", "Digital Marketing", "Google, Meta & SEO"),
        ("◆ BRAND", "Brand & Creative", "Identity, Design & Content"),
    ])

    subsection(doc, "Investment Overview — At a Glance")
    headers = ["Service Line", "Basic", "Standard", "Premium"]
    rows = [
        ["Website (One-Time)", "PKR 60,000", "PKR 120,000", "PKR 250,000"],
        ["Social Media (Monthly)", "PKR 25,000", "PKR 45,000", "PKR 75,000"],
        ["Digital Marketing (Monthly)", "PKR 35,000", "PKR 65,000", "PKR 120,000"],
    ]
    styled_table(doc, headers, rows)

    callout_box(doc, "ENGAGEMENT FLEXIBILITY",
                "Packages may be selected individually or combined. Ad media spend is billed separately. "
                "Pricing valid for 30 days. Detailed inclusions follow in each dedicated service section.")

    add_page_break(doc)


def build_website_section(doc):
    section_header(doc, "04  |  Website Designing & Development",
                   "Premium web solutions engineered for performance and growth")

    body(doc,
         "Your website is the digital headquarters of your business. CoLab Space Point designs and "
         "develops websites that communicate credibility, convert visitors into customers, and scale "
         "alongside your ambitions. WordPress is one of our primary development platforms, enabling "
         "flexible content management and long-term maintainability. We also deliver complete custom "
         "website designing and development solutions tailored precisely to client requirements, "
         "industry standards, and commercial objectives.")

    subsection(doc, "What We Deliver")
    for item in [
        "Strategic information architecture and conversion-oriented page structures",
        "Modern, responsive UI/UX aligned with your brand identity",
        "WordPress CMS implementation with clean, secure configurations",
        "E-commerce capabilities via WooCommerce where required",
        "SEO foundations, analytics, and marketing pixel integrations",
        "Performance optimization, SSL, and security hardening",
        "Admin training and structured post-launch support",
    ]:
        check_item(doc, item)

    add_para(doc, "", space_after=6)
    # ── All three packages as premium pricing cards ──
    add_page_break(doc)
    subsection(doc, "Website Packages — Basic · Standard · Premium")

    packages = [
        {
            "name": "Basic",
            "price": "PKR 60,000",
            "suitable": "Small Businesses & Startups",
            "features": [
                "Up to 5 Pages",
                "Responsive Design",
                "WordPress CMS + Custom UI",
                "Contact Form",
                "WhatsApp Integration",
                "Social Media Integration",
                "Basic SEO + Analytics",
                "Speed + SSL + Security",
                "Admin Training",
                "30 Days Support",
                "⚠ No E-Commerce",
            ],
        },
        {
            "name": "Standard",
            "price": "PKR 120,000",
            "suitable": "Growing Businesses & Stores",
            "features": [
                "Up to 10 Pages",
                "Premium UI/UX + Blog",
                "WooCommerce Store",
                "Product Upload",
                "Payment Gateway",
                "Search Console",
                "Facebook Pixel",
                "Advanced SEO",
                "Speed Optimization",
                "Admin Training",
                "60 Days Support",
            ],
        },
        {
            "name": "Premium",
            "price": "PKR 250,000",
            "suitable": "Scalable Brand Platforms",
            "features": [
                "Unlimited Pages",
                "Fully Custom Design",
                "Advanced WooCommerce",
                "Unlimited Products",
                "Payment Gateway",
                "CRM Integration",
                "Booking System",
                "API Integration",
                "Advanced SEO",
                "Premium Security",
                "90 Days Support",
            ],
        },
    ]
    three_col_packages(doc, packages)

    callout_box(doc, "PLATFORM NOTE",
                "WordPress is one of our primary development platforms, enabling flexible content "
                "management and long-term maintainability. We also provide complete custom website "
                "designing and development solutions according to client requirements.")

    add_page_break(doc)

    # Detailed package narratives
    subsection(doc, "Package Detail — Basic Website  |  PKR 60,000")
    body(doc,
         "The Basic Website package establishes a professional digital presence for small businesses "
         "and startups that need credibility, clarity, and essential functionality. It includes a "
         "custom UI WordPress site of up to five pages, responsive design, contact and WhatsApp "
         "integrations, basic SEO, Google Analytics, speed optimization, SSL, security setup, "
         "admin training, and 30 days of post-launch support.")
    add_para(doc, "⚠  E-Commerce / Online Store is NOT included in the Basic package.",
             size=10, bold=True, color=RGBColor(0x8B, 0x45, 0x13), space_before=4, space_after=8)

    subsection(doc, "Package Detail — Standard Website  |  PKR 120,000")
    body(doc,
         "The Standard Website package is engineered for businesses ready to sell online. It expands "
         "to ten pages with premium UI/UX, a blog module, WooCommerce store setup, product upload, "
         "payment gateway integration, Google Search Console, Facebook Pixel, advanced on-page SEO, "
         "speed optimization, admin training, and 60 days of post-launch support.")

    subsection(doc, "Package Detail — Premium Website  |  PKR 250,000")
    body(doc,
         "The Premium Website package is our flagship offering for brands requiring unlimited scale "
         "and advanced integrations. It includes unlimited pages, a fully custom design system, "
         "advanced WooCommerce with unlimited products, payment gateway, CRM integration, booking "
         "system, custom API integration, advanced SEO architecture, premium security hardening, "
         "enterprise performance optimization, comprehensive admin training, and 90 days of priority support.")

    add_page_break(doc)

    # Comparison Table
    subsection(doc, "Website Packages — Feature Comparison")
    body(doc,
         "Compare packages side by side to identify the solution that best aligns with your "
         "business stage and commercial requirements.")

    headers = ["Feature", "Basic", "Standard", "Premium"]
    rows = [
        ["Investment", "PKR 60,000", "PKR 120,000", "PKR 250,000"],
        ["Pages", "Up to 5", "Up to 10", "Unlimited"],
        ["Design Level", "Custom UI", "Premium UI/UX", "Fully Custom"],
        ["WordPress CMS + Responsive", "✓", "✓", "✓"],
        ["Blog Module", "—", "✓", "✓"],
        ["E-Commerce / WooCommerce", "—", "✓", "Advanced"],
        ["Product Upload", "—", "Included", "Unlimited"],
        ["Payment Gateway", "—", "✓", "✓"],
        ["CRM / Booking / API", "—", "—", "✓"],
        ["Contact + WhatsApp", "✓", "✓", "✓"],
        ["SEO Setup", "Basic", "Advanced", "Advanced"],
        ["Analytics / Search Console / Pixel", "Analytics", "All Three", "All Three"],
        ["Speed + Security + SSL", "✓", "✓", "Premium"],
        ["Admin Training", "✓", "✓", "Comprehensive"],
        ["Post-Launch Support", "30 Days", "60 Days", "90 Days"],
    ]
    styled_table(doc, headers, rows, col_widths=[2.8, 1.4, 1.4, 1.4])

    callout_box(doc, "RECOMMENDATION",
                "Not sure which package fits? Book a free discovery consultation. We will assess "
                "your goals, competitors, and technical requirements to recommend the best path forward.")

    add_page_break(doc)


def build_smm_section(doc):
    section_header(doc, "05  |  Social Media Management",
                   "Monthly packages that build presence, engagement, and brand authority")

    body(doc,
         "Social media is where modern brands earn attention, trust, and conversations. "
         "CoLab Space Point manages your social channels with strategic content calendars, "
         "platform-native creative, community engagement, and clear performance reporting — "
         "so your brand remains consistent, visible, and commercially relevant every month.")

    subsection(doc, "Monthly Social Media Packages")

    packages = [
        {
            "name": "Basic",
            "price": "PKR 25,000/mo",
            "suitable": "Emerging Brands",
            "features": [
                "Facebook Management",
                "Instagram Management",
                "12 Posts Per Month",
                "Professional Captions",
                "Strategic Hashtags",
                "Monthly Performance Report",
                "Content Calendar Planning",
                "Brand-Aligned Visuals",
            ],
        },
        {
            "name": "Standard",
            "price": "PKR 45,000/mo",
            "suitable": "Growing Businesses",
            "features": [
                "Facebook Management",
                "Instagram Management",
                "LinkedIn Management",
                "20 Posts Per Month",
                "Stories Content",
                "Reels Planning",
                "Community Management",
                "Analytics & Insights",
                "Content Calendar",
                "Hashtag & Caption Strategy",
            ],
        },
        {
            "name": "Premium",
            "price": "PKR 75,000/mo",
            "suitable": "Brand Leaders",
            "features": [
                "Facebook Management",
                "Instagram Management",
                "LinkedIn Management",
                "TikTok Management",
                "30+ Posts Per Month",
                "Daily Stories",
                "Reels Strategy & Execution",
                "Full Community Management",
                "Weekly Performance Reports",
                "Competitor Monitoring",
                "Growth Strategy Sessions",
                "Priority Creative Support",
            ],
        },
    ]
    three_col_packages(doc, packages)

    subsection(doc, "Social Media — Package Comparison")
    headers = ["Deliverable", "Basic", "Standard", "Premium"]
    rows = [
        ["Monthly Investment", "PKR 25,000", "PKR 45,000", "PKR 75,000"],
        ["Platforms", "FB + Instagram", "FB + IG + LinkedIn", "+ TikTok"],
        ["Posts Per Month", "12", "20", "30+"],
        ["Captions & Hashtags", "✓", "✓", "✓"],
        ["Stories", "—", "✓", "Daily"],
        ["Reels", "—", "Planning", "Full Strategy"],
        ["Community Management", "—", "✓", "✓"],
        ["Reporting", "Monthly", "Analytics", "Weekly"],
    ]
    styled_table(doc, headers, rows)

    callout_box(doc, "CONTENT QUALITY STANDARD",
                "All packages include professional captions, brand-consistent visuals, and platform-optimized "
                "formatting. Premium adds reels strategy and weekly performance reviews with actionable recommendations.")

    add_page_break(doc)


def build_digital_marketing(doc):
    section_header(doc, "06  |  Digital Marketing",
                   "Performance packages engineered for leads, conversions, and measurable ROI")

    body(doc,
         "Digital marketing at CoLab Space Point is disciplined, transparent, and outcome-oriented. "
         "We plan, launch, and optimize paid and organic campaigns across Meta and Google ecosystems — "
         "with clear tracking frameworks so every rupee of your budget is accountable to business results.")

    subsection(doc, "Monthly Digital Marketing Packages")

    packages = [
        {
            "name": "Basic",
            "price": "PKR 35,000/mo",
            "suitable": "Focused Ad Growth",
            "features": [
                "Meta Ads Management",
                "Audience Targeting",
                "Campaign Optimization",
                "Monthly Reporting",
                "Ad Creative Guidance",
                "Budget Recommendations",
                "A/B Testing Basics",
                "Performance Review Call",
            ],
        },
        {
            "name": "Standard",
            "price": "PKR 65,000/mo",
            "suitable": "Multi-Channel Growth",
            "features": [
                "Google Ads Management",
                "Meta Ads Management",
                "Lead Generation Campaigns",
                "Conversion Tracking Setup",
                "Landing Page Recommendations",
                "Monthly Performance Reports",
                "Audience & Keyword Strategy",
                "Campaign Optimization",
                "Retargeting Foundations",
            ],
        },
        {
            "name": "Premium",
            "price": "PKR 120,000/mo",
            "suitable": "Full-Funnel Growth",
            "features": [
                "Google Ads Management",
                "Meta Ads Management",
                "SEO Strategy & Execution",
                "Remarketing Campaigns",
                "Conversion Rate Optimization",
                "Full Marketing Strategy",
                "Weekly Strategy Meetings",
                "Detailed Analytics Reports",
                "Competitor Analysis",
                "Funnel Optimization",
                "Priority Account Management",
            ],
        },
    ]
    three_col_packages(doc, packages)

    subsection(doc, "Digital Marketing — Package Comparison")
    headers = ["Capability", "Basic", "Standard", "Premium"]
    rows = [
        ["Monthly Investment", "PKR 35,000", "PKR 65,000", "PKR 120,000"],
        ["Meta Ads", "✓", "✓", "✓"],
        ["Google Ads", "—", "✓", "✓"],
        ["SEO", "—", "—", "✓"],
        ["Lead Generation + Tracking", "—", "✓", "✓"],
        ["Remarketing", "—", "Foundations", "✓"],
        ["Landing Page Advice", "—", "✓", "✓"],
        ["Conversion Optimization", "—", "—", "✓"],
        ["Marketing Strategy", "—", "—", "✓"],
        ["Reporting / Meetings", "Monthly", "Monthly", "Weekly + Detailed"],
    ]
    styled_table(doc, headers, rows)

    callout_box(doc, "AD SPEND NOTE",
                "Package fees cover strategy, setup, management, optimization, and reporting. "
                "Media spend (Google & Meta ad budgets) is billed separately and controlled entirely by the client. "
                "We recommend budgets based on industry benchmarks during onboarding.")

    add_page_break(doc)


def build_addons(doc):
    section_header(doc, "07  |  Add-On Services",
                   "Optional enhancements to strengthen and extend your digital ecosystem")

    body(doc,
         "Complement your primary package with specialist services that elevate branding, content, "
         "operations, and technical infrastructure. Add-ons may be engaged independently or "
         "bundled with any website, social, or marketing package.")

    headers = ["Service", "Description", "Starting From"]
    rows = [
        ["Logo Design", "Professional logo concepts with revisions & final files", "PKR 15,000"],
        ["Brand Identity", "Color system, typography, brand guidelines package", "PKR 40,000"],
        ["Landing Page Design", "High-conversion single-page design & development", "PKR 25,000"],
        ["Website Maintenance", "Updates, backups, security & minor content changes (monthly)", "PKR 8,000/mo"],
        ["Content Writing", "SEO-aware website or blog content (per page/article)", "PKR 3,000"],
        ["Graphic Design", "Social creatives, banners, promotional visuals (per design)", "PKR 2,500"],
        ["Video Editing", "Short-form promotional or reels editing (per video)", "PKR 5,000"],
        ["SEO Audit", "Technical & on-page audit with prioritized action plan", "PKR 20,000"],
        ["Product Upload", "Bulk product listing, images & descriptions (per batch)", "PKR 10,000"],
        ["Business Email Setup", "Professional email on your domain (setup & configuration)", "PKR 5,000"],
        ["Domain & Hosting Assistance", "Domain registration guidance & hosting configuration support", "PKR 3,000"],
    ]
    styled_table(doc, headers, rows, col_widths=[2.2, 3.3, 1.5])

    add_para(doc,
             "Complex brand systems or enterprise retainers are scoped individually — contact us for a tailored quotation.",
             size=9.5, color=GRAY, italic=True, space_after=4)

    add_page_break(doc)


def build_process(doc):
    section_header(doc, "08  |  Our Process",
                   "A disciplined seven-stage methodology from discovery to ongoing support")

    body(doc,
         "Every engagement at CoLab Space Point follows a structured delivery framework. "
         "This ensures clarity, quality control, and predictable outcomes — whether we are "
         "launching a website, activating social channels, or scaling paid campaigns.")

    stages = [
        ("01", "DISCOVERY",
         "Deep-dive consultation covering business model, audience, competitors, brand positioning, "
         "and commercial objectives — defining success criteria and project boundaries."),
        ("02", "PLANNING",
         "Roadmap creation including sitemap or campaign structure, timelines, milestones, "
         "resource allocation, and clearly defined deliverables."),
        ("03", "DESIGN",
         "Creative concepts and UI/UX directions reflecting your brand identity. You review and "
         "approve designs before development begins."),
        ("04", "DEVELOPMENT",
         "Engineering of approved designs into high-performance assets — CMS, integrations, "
         "responsive build, content population, or campaign setup and tracking."),
        ("05", "TESTING",
         "Quality assurance across devices, browsers, forms, payment flows, tracking pixels, "
         "and performance benchmarks before any public release."),
        ("06", "LAUNCH",
         "Coordinated go-live with DNS/hosting verification, analytics confirmation, and a "
         "structured launch checklist for a smooth release."),
        ("07", "SUPPORT",
         "Post-launch support, training, reporting, and optimization — spanning 30 to 90 days "
         "by package, with options for ongoing retainers."),
    ]

    for num, title, desc in stages:
        table = doc.add_table(rows=1, cols=2)
        c0, c1 = table.rows[0].cells
        shade_cell(c0, HEADER_FILL)
        set_cell_borders(c0, color=HEADER_FILL, sz="4")
        clear_cell_margins(c0, top=28, bottom=28, left=30, right=30)
        cell_text(c0, f"{num}  {title}", size=9, bold=True, color=WHITE, align="center")
        c0.width = Inches(1.5)

        shade_cell(c1, LIGHT_FILL)
        set_cell_borders(c1, color="D0D8E0", sz="4")
        clear_cell_margins(c1, top=28, bottom=28, left=80, right=50)
        cell_text(c1, desc, size=9.5, color=DARK)
        c1.width = Inches(5.5)

    add_page_break(doc)


def build_trust(doc):
    section_header(doc, "09  |  Why Clients Trust Us",
                   "Partnership built on integrity, delivery excellence, and measurable impact")

    body(doc,
         "Trust is earned through consistency. Since establishing the CoLab Point ecosystem in 2021 "
         "in Gujrat, we have cultivated a community of freelancers, startups, and businesses who "
         "rely on our environment, expertise, and commitment. CoLab Space Point extends that same "
         "standard into every digital engagement we undertake.")

    commitments = [
        ("Clarity Before Commitment",
         "Scopes, timelines, and investments defined before work begins — no hidden fees or mid-project surprises."),
        ("Quality Without Compromise",
         "Every asset is delivered to a professional standard worthy of a competitive marketplace."),
        ("Results Over Vanity Metrics",
         "We prioritize qualified inquiries, sales, brand recognition, and sustainable growth — not empty impressions."),
        ("Local Roots, Global Standards",
         "Based in Gujrat with deep Pakistani market understanding, delivered to international agency standards."),
        ("Partnership Mentality",
         "We succeed when you succeed — invested, attentive, and accountable for the long term."),
    ]

    for title, desc in commitments:
        add_para(doc, f"▸  {title}", size=11, bold=True, color=NAVY, space_before=6, space_after=2)
        body(doc, desc, space_after=2)

    callout_box(doc, "CLIENT PROMISE",
                "When you engage CoLab Space Point, you gain a dedicated digital partner — not a "
                "transactional vendor. We communicate with professionalism, deliver with precision, "
                "and stand behind the quality of every project that carries our name.")

    add_page_break(doc)


def build_faq(doc):
    section_header(doc, "10  |  Frequently Asked Questions",
                   "Clear answers to the questions clients ask most often")

    faqs = [
        ("How long does a website project typically take?",
         "Basic: 2–3 weeks after content/approvals. Standard: 3–5 weeks. Premium: 6–8 weeks."),
        ("Do you use WordPress for all websites?",
         "WordPress is our primary platform. We also deliver custom solutions per client requirements."),
        ("Is advertising budget included in Digital Marketing packages?",
         "No. Package fees cover management and reporting. Media spend is paid separately by the client."),
        ("Can I combine Website, Social Media, and Marketing packages?",
         "Yes. Bundled engagements with coordinated strategy across all channels are available."),
        ("What do I need to provide to get started?",
         "Business info, brand assets, content/products, audience insights, and access credentials."),
        ("Do you provide training and post-launch support?",
         "Yes. All website packages include admin training and 30–90 days support by package level."),
        ("Do you work outside Gujrat, and how do we begin?",
         "Yes — across Pakistan and internationally. Contact us for a discovery consultation to confirm scope."),
    ]

    for i, (q, a) in enumerate(faqs):
        table = doc.add_table(rows=2, cols=1)
        c0 = table.rows[0].cells[0]
        shade_cell(c0, HEADER_FILL if i % 2 == 0 else ACCENT_FILL)
        set_cell_borders(c0, color=HEADER_FILL, sz="4")
        clear_cell_margins(c0, top=18, bottom=18, left=60, right=40)
        cell_text(c0, f"Q.  {q}", size=9, bold=True, color=WHITE)

        c1 = table.rows[1].cells[0]
        shade_cell(c1, LIGHT_FILL)
        set_cell_borders(c1, color="D0D8E0", sz="4")
        clear_cell_margins(c1, top=18, bottom=18, left=60, right=40)
        cell_text(c1, f"A.  {a}", size=8.5, color=DARK)

    add_page_break(doc)


def build_contact(doc):
    section_header(doc, "11  |  Contact & Next Steps",
                   "Let's build your digital advantage together")

    body(doc,
         "Ready to elevate your digital presence? Our team at CoLab Space Point is prepared to "
         "discuss your objectives, recommend the right package combination, and outline a clear "
         "path from kickoff to launch.")

    subsection(doc, "How to Engage")
    steps = [
        "Contact us via phone, email, or website to request a discovery consultation.",
        "Share your business goals, digital assets, and preferred package interests.",
        "Receive a confirmed scope, timeline, and investment summary.",
        "Approve the engagement and begin onboarding with our project team.",
    ]
    for i, s in enumerate(steps, 1):
        add_para(doc, f"{i}.  {s}", size=10, color=DARK, space_before=1, space_after=2)

    add_para(doc, "", space_after=4)

    # Contact card
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, HEADER_FILL)
    set_cell_borders(cell, color=HEADER_FILL, sz="4")
    clear_cell_margins(cell, top=70, bottom=70, left=120, right=120)
    cell.text = ""

    lines = [
        ("COLAB SPACE POINT", 16, True, WHITE, "Georgia"),
        ("Digital Agency  ·  A CoLab Point Initiative", 10, False, RGBColor(0xA0, 0xC0, 0xE0), "Calibri"),
        ("", 4, False, WHITE, "Calibri"),
        ("OFFICE ADDRESS", 8, True, RGBColor(0xC9, 0xA2, 0x27), "Calibri"),
        ("2nd Floor, Anwar Center, Madina Road, Near Gymkhana, Gujrat, Pakistan", 10, False, WHITE, "Calibri"),
        ("", 4, False, WHITE, "Calibri"),
        ("PHONE", 8, True, RGBColor(0xC9, 0xA2, 0x27), "Calibri"),
        ("+92 349 7684322  ·  +92 332 4384322", 11, True, WHITE, "Calibri"),
        ("", 4, False, WHITE, "Calibri"),
        ("EMAIL", 8, True, RGBColor(0xC9, 0xA2, 0x27), "Calibri"),
        ("colabpoint@gmail.com  ·  hello@colabpoint.com", 10, False, WHITE, "Calibri"),
        ("", 4, False, WHITE, "Calibri"),
        ("WEBSITE", 8, True, RGBColor(0xC9, 0xA2, 0x27), "Calibri"),
        ("www.colabpoint.com", 11, True, WHITE, "Calibri"),
    ]

    first = True
    for text, size, bold, color, font in lines:
        if first:
            p = cell.paragraphs[0]
            first = False
        else:
            p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        if text:
            r = p.add_run(text)
            set_run(r, size=size, bold=bold, color=color, font=font)

    add_para(doc, "", space_after=8)

    # Closing statement
    add_para(doc, "Thank You", size=16, bold=True, color=NAVY,
             align="center", font="Georgia", space_before=4, space_after=4)
    add_para(doc,
             "We appreciate the opportunity to present this proposal and look forward to "
             "partnering with you on your digital growth journey.",
             size=10, color=GRAY, align="center", space_after=6)

    add_thin_line(doc)

    add_para(doc,
             "© CoLab Space Point  ·  CoLab Point  ·  Gujrat, Pakistan  ·  All Rights Reserved",
             size=8, color=GRAY, align="center", space_before=4, space_after=1)
    add_para(doc,
             "This document is confidential and intended solely for the recipient. "
             "Pricing valid for 30 days from the date of issue unless otherwise stated.",
             size=8, color=GRAY, align="center", italic=True, space_after=2)


def build_executive_summary(doc):
    """Premium overview page."""
    section_header(doc, "Executive Summary",
                   "A concise overview of the opportunity and our recommended approach")

    body(doc,
         "In today's competitive marketplace, a fragmented digital presence is a liability. "
         "Businesses that treat their website, social channels, and paid media as isolated "
         "activities forfeit the compounding advantage of an integrated digital strategy. "
         "CoLab Space Point exists to close that gap — with clear, investment-transparent packages "
         "so you can select the precise engagement level your business requires today, with a "
         "path to scale tomorrow.")

    subsection(doc, "Proposal Highlights")
    highlights = [
        "Three website packages from PKR 60,000 to PKR 250,000 — startup sites through advanced e-commerce",
        "Three social media retainers from PKR 25,000 to PKR 75,000 per month",
        "Three digital marketing packages from PKR 35,000 to PKR 120,000 per month",
        "Eleven optional add-ons for branding, content, maintenance, and technical support",
        "Seven-stage delivery process from Discovery through ongoing Support",
        "Transparent communication, defined support windows, and partnership-oriented engagement",
    ]
    for h in highlights:
        check_item(doc, h)

    callout_box(doc, "NEXT STEP",
                "Review the service sections that follow, identify packages aligned with your priorities, "
                "and contact our team for a complimentary discovery consultation tailored to your industry and goals.")

    add_page_break(doc)


# ── Main ──────────────────────────────────────────────────────

def main():
    doc = Document()
    set_narrow_margins(doc.sections[0])
    add_header(doc)
    add_footer(doc)

    # Style default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = DARK

    build_cover(doc)
    build_toc(doc)
    build_executive_summary(doc)
    build_company_profile(doc)
    build_why_choose(doc)
    build_services_overview(doc)
    build_website_section(doc)
    build_smm_section(doc)
    build_digital_marketing(doc)
    build_addons(doc)
    build_process(doc)
    build_trust(doc)
    build_faq(doc)
    build_contact(doc)

    doc.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
