#!/usr/bin/env python3
"""CoLab Point proposal — Word + PDF with embedded photos."""

import io
import urllib.request
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fpdf import FPDF
from PIL import Image

ASSETS = Path("/workspace/assets/proposal")
OUT_DOCX = "/workspace/CoLab_Space_Point_Proposal_Photos.docx"
OUT_PDF = "/workspace/CoLab_Space_Point_Proposal_Photos.pdf"
# Keep legacy name too
OUT_DOCX_ALT = "/workspace/CoLab_Space_Point_Digital_Agency_Proposal.docx"

BLACK = RGBColor(0, 0, 0)
TEAL = RGBColor(0x06, 0xAC, 0xBA)
NAVY = RGBColor(0x04, 0x24, 0x3C)
FONT = "Arial"

COMPANY = "CoLab Space Point"
WEB = "www.colabpoint.com"
PHONE = "+92 349 7684322  |  +92 478 986460"
EMAIL = "colabpoint@gmail.com  |  hello@colabpoint.com"
ADDRESS = "2nd Floor Anwar Center, Madina Road Near Gymkhana, Gujrat, Pakistan"

IMAGES = {
    "logo": "https://colabpoint.com/wp-content/uploads/2024/05/Web-Logo-II.png",
    "cover_banner": "https://colabpoint.com/wp-content/uploads/2024/06/popup-bg-1-1024x374-1.jpg",
    "office": "https://images.unsplash.com/photo-1497366216548-37526070297c?w=1200&q=80",
    "web_design": "https://images.unsplash.com/photo-1460925895917-afdab827c52f?w=1200&q=80",
    "ecommerce": "https://images.unsplash.com/photo-1556742049-0cfed4f6a45d?w=1200&q=80",
    "premium_web": "https://images.unsplash.com/photo-1498050108023-c5249f4df085?w=1200&q=80",
    "digital_marketing": "https://images.unsplash.com/photo-1551288049-bebda4e38f71?w=1200&q=80",
    "contact": "https://images.unsplash.com/photo-1522071820081-009f0129c71c?w=1200&q=80",
}


def download_all_images():
    ASSETS.mkdir(parents=True, exist_ok=True)
    paths = {}
    for key, url in IMAGES.items():
        for ext in (".jpg", ".jpeg", ".png", ".webp"):
            cached = ASSETS / f"{key}{ext}"
            if cached.exists() and cached.stat().st_size > 1000:
                paths[key] = cached
                break
        else:
            try:
                req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
                data = urllib.request.urlopen(req, timeout=25).read()
                ext = ".jpg" if url.lower().endswith((".jpg", ".jpeg")) else ".png"
                dest = ASSETS / f"{key}{ext}"
                dest.write_bytes(data)
                paths[key] = dest
            except Exception as e:
                print(f"Warning: could not download {key}: {e}")
    return paths


def to_jpeg(path: Path) -> Path:
    """Convert any image to JPEG for reliable PDF embedding."""
    jpg = ASSETS / f"{path.stem}_pdf.jpg"
    if jpg.exists() and jpg.stat().st_mtime >= path.stat().st_mtime:
        return jpg
    img = Image.open(path).convert("RGB")
    img.save(jpg, "JPEG", quality=88)
    return jpg


def force_white_page(doc):
    bg = OxmlElement("w:background")
    bg.set(qn("w:color"), "FFFFFF")
    doc.element.insert(0, bg)


def write_run(paragraph, text, size=14, bold=False, color=BLACK):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FONT
    run.font.color.rgb = color


def add_photo(doc, path: Path, width=6.5):
    if not path or not path.exists():
        return False
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(12)
    p.add_run().add_picture(str(path), width=Inches(width))
    return True


def line(doc, text="", size=14, bold=False, color=BLACK, center=False, space=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        write_run(p, text, size, bold, color)


def section_title(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    write_run(p, text, 20, True, BLACK)


def bullet(doc, text):
    line(doc, f"  •  {text}", 13, False, BLACK, space=4)


def build_docx(img_paths: dict):
    doc = Document()
    force_white_page(doc)
    sec = doc.sections[0]
    for m in (sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin):
        pass
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(2)
    sec.left_margin = Cm(2)
    sec.right_margin = Cm(2)

    # Cover
    add_photo(doc, img_paths.get("cover_banner"), 6.5)
    add_photo(doc, img_paths.get("logo"), 2.0)
    line(doc, "CoLab Point", 16, True, TEAL, center=True)
    line(doc, COMPANY, 28, True, BLACK, center=True)
    line(doc, "Digital Agency Proposal", 18, True, BLACK, center=True)
    line(doc, "Website Designing & Development  |  Digital Marketing", 13, False, BLACK, center=True)
    line(doc, WEB, 13, True, NAVY, center=True)
    line(doc, PHONE, 12, False, BLACK, center=True)
    line(doc, EMAIL, 12, False, BLACK, center=True)
    line(doc, ADDRESS, 12, False, BLACK, center=True)

    doc.add_page_break()
    section_title(doc, "Company Profile")
    add_photo(doc, img_paths.get("office"), 6.5)
    line(
        doc,
        f"{COMPANY} is the digital services arm of CoLab Point, Gujrat. We build websites, "
        "e-commerce stores, and run performance marketing campaigns.",
    )

    section_title(doc, "Website Designing & Development")
    add_photo(doc, img_paths.get("web_design"), 6.5)

    packages = [
        ("BASIC WEBSITE — PKR 50,000", "web_design", "Small businesses", [
            "Up to 5 pages, WordPress, custom UI", "Contact form, WhatsApp, Analytics, SSL",
            "30 days support", "SEO & e-commerce NOT included",
        ]),
        ("STANDARD WEBSITE — PKR 80,000", "ecommerce", "Growing brands", [
            "Everything in Basic, plus 10 pages", "WooCommerce, payment gateway, blog",
            "Advanced SEO, Facebook Pixel, 60 days support",
        ]),
        ("PREMIUM WEBSITE — PKR 120,000", "premium_web", "Enterprises", [
            "Everything in Standard, plus custom design", "CRM, booking, API integrations",
            "Premium security, admin training, 90 days support",
        ]),
    ]
    for title, img_key, audience, items in packages:
        doc.add_page_break()
        section_title(doc, title)
        add_photo(doc, img_paths.get(img_key), 6.0)
        line(doc, f"Best for: {audience}", 13, True, BLACK)
        line(doc, "WHAT IS INCLUDED IN THIS PRICE:", 14, True, TEAL)
        for it in items:
            bullet(doc, it)

    doc.add_page_break()
    section_title(doc, "Digital Marketing")
    add_photo(doc, img_paths.get("digital_marketing"), 6.5)
    for name, price, items in [
        ("BASIC", "15,000", ["Meta Ads", "Audience targeting", "Optimization", "Monthly report"]),
        ("STANDARD", "30,000", ["Google + Meta Ads", "Lead generation", "Conversion tracking", "Monthly reports"]),
        ("PREMIUM", "50,000", ["Google + Meta Ads", "SEO", "Remarketing", "Weekly meetings", "Detailed reports"]),
    ]:
        line(doc, f"{name} — PKR {price} / month", 15, True, BLACK)
        line(doc, "What is included:", 13, True, TEAL)
        for it in items:
            bullet(doc, it)
        line(doc, "")

    doc.add_page_break()
    section_title(doc, "Contact")
    add_photo(doc, img_paths.get("contact"), 5.5)
    line(doc, f"Company: {COMPANY}")
    line(doc, f"Website: {WEB}")
    line(doc, f"Phone: {PHONE}")
    line(doc, f"Email: {EMAIL}")
    line(doc, f"Address: {ADDRESS}")

    doc.save(OUT_DOCX)
    doc.save(OUT_DOCX_ALT)
    print(f"Word saved: {OUT_DOCX} ({Path(OUT_DOCX).stat().st_size // 1024} KB)")


class ProposalPDF(FPDF):
    def header(self):
        pass

    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(100, 100, 100)
        self.cell(0, 8, f"{COMPANY} | {WEB}", align="C")


def pdf_image(pdf, path: Path, w=190):
    if path and path.exists():
        jpg = to_jpeg(path)
        pdf.image(str(jpg), x=10, w=w)
        pdf.ln(4)


def build_pdf(img_paths: dict):
    pdf = ProposalPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf_image(pdf, img_paths.get("cover_banner"))
    if img_paths.get("logo"):
        pdf.image(str(to_jpeg(img_paths["logo"])), x=85, w=40)
        pdf.ln(8)
    pdf.set_font("Helvetica", "B", 24)
    pdf.set_text_color(4, 36, 60)
    pdf.cell(0, 12, COMPANY, ln=True, align="C")
    pdf.set_font("Helvetica", "", 14)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(0, 8, "Digital Agency Proposal", ln=True, align="C")
    pdf.cell(0, 8, WEB, ln=True, align="C")
    pdf.cell(0, 8, PHONE, ln=True, align="C")

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "Company Profile", ln=True)
    pdf_image(pdf, img_paths.get("office"))
    pdf.set_font("Helvetica", "", 11)
    pdf.multi_cell(0, 6, f"{COMPANY} delivers websites, e-commerce, and digital marketing from Gujrat, Pakistan.")

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "Website Packages", ln=True)
    pdf_image(pdf, img_paths.get("web_design"))
    for title, img_key, items in [
        ("Basic - PKR 50,000", "web_design", "5 pages, WordPress, no SEO/e-commerce"),
        ("Standard - PKR 80,000", "ecommerce", "10 pages, WooCommerce, advanced SEO"),
        ("Premium - PKR 120,000", "premium_web", "Custom design, CRM, booking, API"),
    ]:
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 13)
        pdf.cell(0, 8, title, ln=True)
        pdf_image(pdf, img_paths.get(img_key), w=160)
        pdf.set_font("Helvetica", "", 11)
        pdf.multi_cell(0, 6, items)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "Digital Marketing", new_x="LMARGIN", new_y="NEXT")
    pdf_image(pdf, img_paths.get("digital_marketing"))
    pdf.set_font("Helvetica", "", 11)
    pdf.set_x(10)
    for line_txt in [
        "Basic - PKR 15,000/month: Meta Ads, targeting, optimization",
        "Standard - PKR 30,000/month: Google + Meta, leads, tracking",
        "Premium - PKR 50,000/month: Full funnel, SEO, weekly meetings",
    ]:
        pdf.multi_cell(190, 7, line_txt)
        pdf.ln(2)

    pdf.add_page()
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "Contact", new_x="LMARGIN", new_y="NEXT")
    pdf_image(pdf, img_paths.get("contact"), w=150)
    pdf.set_font("Helvetica", "", 11)
    pdf.set_x(10)
    pdf.multi_cell(190, 7, f"{ADDRESS}\n{PHONE}\n{EMAIL}")

    pdf.output(OUT_PDF)
    print(f"PDF saved: {OUT_PDF} ({Path(OUT_PDF).stat().st_size // 1024} KB)")


def main():
    print("Downloading images...")
    paths = download_all_images()
    print(f"Images ready: {len(paths)} -> {list(paths.keys())}")
    build_docx(paths)
    build_pdf(paths)


if __name__ == "__main__":
    main()
